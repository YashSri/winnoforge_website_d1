#!/usr/bin/env python3
"""
auto_update_changelogs.py
Automated synchronization script for FORGE website change records.
Updates:
  1. CHANGES_LOG.txt
  2. changes.txt
  3. FORGE_Website_Project_Change_Report.xlsx
  4. FORGE_Website_Project_Change_Report.csv
"""

import os
import subprocess
import shutil
from datetime import datetime

ROOT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
CHANGES_LOG = os.path.join(ROOT_DIR, "CHANGES_LOG.txt")
CHANGES_TXT = os.path.join(ROOT_DIR, "changes.txt")
EXCEL_REPORT = os.path.join(ROOT_DIR, "FORGE_Website_Project_Change_Report.xlsx")
CSV_REPORT = os.path.join(ROOT_DIR, "FORGE_Website_Project_Change_Report.csv")

def run_git(cmd):
    res = subprocess.run(cmd, cwd=ROOT_DIR, capture_output=True, text=True)
    return res.stdout.strip()

def get_all_commits():
    raw = run_git(["git", "log", "--reverse", "--pretty=format:%h|||%ad|||%an|||%s", "--date=format:%Y-%m-%d %H:%M:%S %z"])
    commits = []
    for line in raw.split("\n"):
        if not line.strip():
            continue
        parts = line.split("|||")
        if len(parts) == 4:
            h, dt, author, subj = parts
            files_raw = run_git(["git", "show", "--stat", "--oneline", h])
            stat_lines = files_raw.split("\n")[1:-1]
            files_summary = ", ".join([l.strip().split("|")[0].strip() for l in stat_lines[:5]])
            if len(stat_lines) > 5:
                files_summary += f" (+{len(stat_lines)-5} more)"
            commits.append({
                "commit": h,
                "date": dt,
                "author": author,
                "subject": subj,
                "files": files_summary or "Repository root"
            })
    return commits

def get_existing_recorded_commits():
    recorded = set()
    if os.path.exists(CHANGES_LOG):
        with open(CHANGES_LOG, "r", encoding="utf-8") as f:
            for line in f:
                if "COMMIT:" in line:
                    parts = line.split("COMMIT:")
                    if len(parts) > 1:
                        c = parts[1].split("|")[0].strip()
                        recorded.add(c)
    return recorded

def sync():
    commits = get_all_commits()
    recorded = get_existing_recorded_commits()
    new_commits = [c for c in commits if c["commit"] not in recorded]
    
    print(f"Total Git Commits: {len(commits)} | Already Recorded: {len(recorded)} | New: {len(new_commits)}")
    
    if not new_commits:
        print("All change logs are currently up to date.")
        return False
        
    # Append to CHANGES_LOG.txt before the footer
    with open(CHANGES_LOG, "r", encoding="utf-8") as f:
        content = f.read()
        
    footer_marker = "================================================================================\n4. CURRENT PRODUCTION STATUS & VERIFICATION SUMMARY"
    parts = content.split(footer_marker)
    
    body = parts[0]
    footer = footer_marker + (parts[1] if len(parts) > 1 else "")
    
    current_count = len(recorded)
    new_entries_text = ""
    for c in new_commits:
        current_count += 1
        new_entries_text += f"""
--------------------------------------------------------------------------------
RECORD #{current_count:02d} | COMMIT: {c['commit']} | TIMESTAMP: {c['date']}
--------------------------------------------------------------------------------
AUTHOR / COMMITTED BY:
  {c['author']}

USER INSTRUCTION / REQUEST:
  {c['subject']}

SITE CHANGES IMPLEMENTED:
  {c['subject']}

KEY FILES AFFECTED:
  {c['files']}
"""
    
    updated_content = body.rstrip() + "\n" + new_entries_text + "\n" + footer
    with open(CHANGES_LOG, "w", encoding="utf-8") as f:
        f.write(updated_content)
        
    shutil.copyfile(CHANGES_LOG, CHANGES_TXT)
    print(f"Appended {len(new_commits)} new entries to {CHANGES_LOG} and {CHANGES_TXT}")
    
    # Update Excel
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        
        wb = openpyxl.load_workbook(EXCEL_REPORT)
        ws_log = wb["Detailed Changelog"]
        ws_summary = wb["Executive Summary"]
        
        thin_side = Side(style="thin", color="D1D5DB")
        border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        font_regular = Font(name="Calibri", size=10, color="0F172A")
        font_bold = Font(name="Calibri", size=10, bold=True, color="0F172A")
        font_mono = Font(name="Consolas", size=9.5, color="1E293B")
        font_pill_green = Font(name="Calibri", size=9.5, bold=True, color="059669")
        fill_light_green = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
        fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        
        align_wrap_center = Alignment(horizontal="center", vertical="top", wrap_text=True)
        align_wrap_left = Alignment(horizontal="left", vertical="top", wrap_text=True)
        
        start_row = ws_log.max_row + 1
        for idx, c in enumerate(new_commits, start_row - 1):
            row_idx = ws_log.max_row + 1
            ws_log.row_dimensions[row_idx].height = 40
            
            cells = [
                (1, f"#{idx:02d}", align_wrap_center, font_bold),
                (2, c["date"], align_wrap_center, font_mono),
                (3, "Continuous Updates", align_wrap_left, font_regular),
                (4, "Update", align_wrap_center, font_bold),
                (5, c["commit"], align_wrap_center, font_mono),
                (6, c["author"], align_wrap_center, font_regular),
                (7, c["subject"], align_wrap_left, font_regular),
                (8, c["subject"], align_wrap_left, font_regular),
                (9, c["files"], align_wrap_left, font_mono),
                (10, "Committed", align_wrap_center, font_pill_green)
            ]
            for col_idx, val, alignment, font_style in cells:
                cell = ws_log.cell(row=row_idx, column=col_idx, value=val)
                cell.font = font_style
                cell.alignment = alignment
                cell.border = border_cell
                if col_idx == 10:
                    cell.fill = fill_light_green
                elif row_idx % 2 == 1:
                    cell.fill = fill_zebra
                    
        # Update KPI in summary
        ws_summary.cell(row=7, column=2, value=f"{len(commits)} Commits Recorded")
        ws_summary.cell(row=8, column=2, value=f"{len(commits)} Actionable Milestones")
        
        wb.save(EXCEL_REPORT)
        print(f"Updated Excel workbook: {EXCEL_REPORT}")
        
        # Export CSV
        import csv
        with open(CSV_REPORT, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            for row in ws_log.iter_rows(values_only=True):
                writer.writerow(row)
        print(f"Updated CSV export: {CSV_REPORT}")
        
    except Exception as e:
        print("Excel sync error:", e)
        
    return True

if __name__ == "__main__":
    sync()
